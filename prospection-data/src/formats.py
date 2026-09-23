from __future__ import annotations

import bz2
import csv
import gzip
import io
import json
import lzma
import zipfile
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, BinaryIO, Iterator
from xml.etree import ElementTree as ET

try:
    import ijson
except Exception:
    ijson = None

@dataclass
class ParsedRecord:
    payload: Any | None
    location: str
    raw_fragment: str | None = None
    error: str | None = None

SUPPORTED_EXTENSIONS = {
    ".csv", ".tsv", ".psv", ".txt", ".json", ".jsonl", ".ndjson",
    ".xml", ".xlsx", ".xls", ".html", ".htm", ".parquet",
}

def _logical_suffix(name: str) -> str:
    suffixes = [s.lower() for s in Path(name).suffixes]
    while suffixes and suffixes[-1] in {".gz", ".bz2", ".xz", ".lzma"}:
        suffixes.pop()
    return suffixes[-1] if suffixes else ""

def infer_format(name: str, explicit: str = "auto") -> str:
    if explicit and explicit != "auto":
        return explicit.lower()
    ext = _logical_suffix(name)
    return {
        ".csv": "csv", ".tsv": "tsv", ".psv": "psv", ".txt": "csv",
        ".json": "json", ".jsonl": "jsonl", ".ndjson": "jsonl",
        ".xml": "xml", ".xlsx": "xlsx", ".xls": "xls",
        ".html": "html", ".htm": "html", ".parquet": "parquet",
    }.get(ext, "jsonl")

def is_supported_path(path: Path) -> bool:
    suffixes = [s.lower() for s in path.suffixes]
    if not suffixes:
        return False
    if suffixes[-1] == ".zip":
        return True
    while suffixes and suffixes[-1] in {".gz", ".bz2", ".xz", ".lzma"}:
        suffixes.pop()
    return bool(suffixes and suffixes[-1] in SUPPORTED_EXTENSIONS)

def _open_decompressed(path: Path) -> BinaryIO:
    suffix = path.suffix.lower()
    if suffix == ".gz":
        return gzip.open(path, "rb")
    if suffix == ".bz2":
        return bz2.open(path, "rb")
    if suffix in {".xz", ".lzma"}:
        return lzma.open(path, "rb")
    return path.open("rb")

def _delimiter(name: str, header: str, explicit: str | None = None) -> str:
    if explicit:
        return explicit
    ext = _logical_suffix(name)
    if ext == ".tsv":
        return "\t"
    if ext == ".psv":
        return "|"
    candidates = [",", ";", "\t", "|"]
    return max(candidates, key=lambda d: header.count(d))

def _prepend(first: str, iterator):
    yield first
    yield from iterator

def _iter_csv(binary: BinaryIO, name: str, delimiter: str | None) -> Iterator[ParsedRecord]:
    text = io.TextIOWrapper(binary, encoding="utf-8-sig", errors="replace", newline="")
    header = text.readline()
    if not header:
        return
    reader = csv.DictReader(_prepend(header, text), delimiter=_delimiter(name, header, delimiter))
    for line_no, row in enumerate(reader, start=2):
        yield ParsedRecord(dict(row), f"{name}:row:{line_no}")

def _iter_jsonl(binary: BinaryIO, name: str) -> Iterator[ParsedRecord]:
    text = io.TextIOWrapper(binary, encoding="utf-8-sig", errors="replace")
    for line_no, line in enumerate(text, start=1):
        raw = line.strip()
        if not raw:
            continue
        try:
            yield ParsedRecord(json.loads(raw), f"{name}:line:{line_no}")
        except Exception as exc:
            yield ParsedRecord(None, f"{name}:line:{line_no}", raw[:65536], str(exc))

def _path_get(data: Any, path: str | None) -> Any:
    if not path:
        return data
    current = data
    for part in path.split("."):
        if not part:
            continue
        if isinstance(current, dict):
            current = current.get(part)
        else:
            return None
    return current

def _first_non_ws(binary: BinaryIO) -> bytes:
    try:
        pos = binary.tell()
        chunk = binary.read(4096)
        binary.seek(pos)
    except Exception:
        return b""
    return chunk.lstrip()[:1]

def _iter_json(binary: BinaryIO, name: str, records_path: str | None) -> Iterator[ParsedRecord]:
    first = _first_non_ws(binary)
    if ijson is not None and (records_path or first == b"["):
        prefix = "item" if not records_path else records_path + ".item"
        try:
            emitted = 0
            for index, item in enumerate(ijson.items(binary, prefix), start=1):
                emitted += 1
                yield ParsedRecord(item, f"{name}:item:{index}")
            if emitted or first == b"[":
                return
            try:
                binary.seek(0)
            except Exception:
                return
        except Exception:
            try:
                binary.seek(0)
            except Exception:
                return
    try:
        data = json.load(io.TextIOWrapper(binary, encoding="utf-8-sig", errors="replace"))
        records = _path_get(data, records_path)
        if isinstance(records, list):
            for index, item in enumerate(records, start=1):
                yield ParsedRecord(item, f"{name}:item:{index}")
        elif records is not None:
            yield ParsedRecord(records, f"{name}:item:1")
    except Exception as exc:
        yield ParsedRecord(None, f"{name}:json", None, str(exc))

def _xml_to_obj(element: ET.Element) -> Any:
    children = list(element)
    if not children:
        return (element.text or "").strip()
    result: dict[str, Any] = {}
    for child in children:
        value = _xml_to_obj(child)
        key = child.tag.split("}")[-1]
        if key in result:
            if not isinstance(result[key], list):
                result[key] = [result[key]]
            result[key].append(value)
        else:
            result[key] = value
    if element.attrib:
        result["_attributes"] = dict(element.attrib)
    return result

def _iter_xml(binary: BinaryIO, name: str, record_tag: str | None) -> Iterator[ParsedRecord]:
    depth = 0
    index = 0
    try:
        for event, elem in ET.iterparse(binary, events=("start", "end")):
            if event == "start":
                depth += 1
                continue
            tag = elem.tag.split("}")[-1]
            should_emit = (record_tag and tag == record_tag) or (not record_tag and depth == 2)
            if should_emit:
                index += 1
                yield ParsedRecord(_xml_to_obj(elem), f"{name}:xml:{index}")
                elem.clear()
            depth -= 1
    except Exception as exc:
        yield ParsedRecord(None, f"{name}:xml", None, str(exc))

class _TableParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.in_cell = False
        self.current_cell: list[str] = []
        self.current_row: list[str] = []
        self.rows: list[list[str]] = []
    def handle_starttag(self, tag, attrs):
        if tag.lower() in {"td", "th"}:
            self.in_cell = True
            self.current_cell = []
    def handle_data(self, data):
        if self.in_cell:
            self.current_cell.append(data)
    def handle_endtag(self, tag):
        tag = tag.lower()
        if tag in {"td", "th"} and self.in_cell:
            self.current_row.append("".join(self.current_cell).strip())
            self.in_cell = False
        elif tag == "tr" and self.current_row:
            self.rows.append(self.current_row)
            self.current_row = []

def _iter_html(binary: BinaryIO, name: str) -> Iterator[ParsedRecord]:
    parser = _TableParser()
    parser.feed(binary.read().decode("utf-8", errors="replace"))
    if not parser.rows:
        return
    headers = parser.rows[0]
    for index, row in enumerate(parser.rows[1:], start=2):
        padded = row + [""] * max(0, len(headers) - len(row))
        yield ParsedRecord(dict(zip(headers, padded)), f"{name}:table-row:{index}")

def _iter_xlsx(binary: BinaryIO, name: str, sheet: str | None) -> Iterator[ParsedRecord]:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(binary.read()), read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        headers = [str(x) if x is not None else "" for x in next(rows, [])]
        for index, values in enumerate(rows, start=2):
            yield ParsedRecord(
                {headers[i]: values[i] if i < len(values) else None for i in range(len(headers))},
                f"{name}:sheet:{ws.title}:row:{index}",
            )
    except Exception as exc:
        yield ParsedRecord(None, f"{name}:xlsx", None, str(exc))

def _iter_xls(binary: BinaryIO, name: str, sheet: str | None) -> Iterator[ParsedRecord]:
    try:
        import xlrd
        book = xlrd.open_workbook(file_contents=binary.read())
        ws = book.sheet_by_name(sheet) if sheet else book.sheet_by_index(0)
        headers = [str(ws.cell_value(0, c)) for c in range(ws.ncols)] if ws.nrows else []
        for r in range(1, ws.nrows):
            yield ParsedRecord(
                {headers[c]: ws.cell_value(r, c) for c in range(ws.ncols)},
                f"{name}:sheet:{ws.name}:row:{r+1}",
            )
    except Exception as exc:
        yield ParsedRecord(None, f"{name}:xls", None, str(exc))

def _iter_parquet(binary: BinaryIO, name: str) -> Iterator[ParsedRecord]:
    try:
        import pyarrow.parquet as pq
    except Exception as exc:
        yield ParsedRecord(None, f"{name}:parquet", None, f"pyarrow required: {exc}")
        return
    try:
        table = pq.read_table(io.BytesIO(binary.read()))
        index = 0
        for batch in table.to_batches(max_chunksize=10000):
            for row in batch.to_pylist():
                index += 1
                yield ParsedRecord(row, f"{name}:row:{index}")
    except Exception as exc:
        yield ParsedRecord(None, f"{name}:parquet", None, str(exc))

def iter_stream_records(
    binary: BinaryIO,
    name: str,
    fmt: str = "auto",
    delimiter: str | None = None,
    records_path: str | None = None,
    xml_record_tag: str | None = None,
    sheet: str | None = None,
) -> Iterator[ParsedRecord]:
    detected = infer_format(name, fmt)
    if detected in {"csv", "tsv", "psv"}:
        yield from _iter_csv(binary, name, delimiter)
    elif detected == "jsonl":
        yield from _iter_jsonl(binary, name)
    elif detected == "json":
        yield from _iter_json(binary, name, records_path)
    elif detected == "xml":
        yield from _iter_xml(binary, name, xml_record_tag)
    elif detected == "xlsx":
        yield from _iter_xlsx(binary, name, sheet)
    elif detected == "xls":
        yield from _iter_xls(binary, name, sheet)
    elif detected == "html":
        yield from _iter_html(binary, name)
    elif detected == "parquet":
        yield from _iter_parquet(binary, name)
    else:
        yield ParsedRecord(None, name, None, f"Unsupported format: {detected}")

def iter_file_records(
    path: Path,
    fmt: str = "auto",
    delimiter: str | None = None,
    records_path: str | None = None,
    xml_record_tag: str | None = None,
    sheet: str | None = None,
) -> Iterator[ParsedRecord]:
    if path.suffix.lower() == ".zip":
        with zipfile.ZipFile(path) as archive:
            for member in archive.infolist():
                if member.is_dir():
                    continue
                member_path = Path(member.filename)
                if not is_supported_path(member_path):
                    continue
                with archive.open(member, "r") as fh:
                    if member_path.suffix.lower() in {".xlsx", ".xls", ".parquet"}:
                        yield from iter_stream_records(
                            io.BytesIO(fh.read()), member.filename, fmt, delimiter,
                            records_path, xml_record_tag, sheet
                        )
                    else:
                        yield from iter_stream_records(
                            fh, member.filename, fmt, delimiter,
                            records_path, xml_record_tag, sheet
                        )
        return
    with _open_decompressed(path) as binary:
        yield from iter_stream_records(
            binary, path.name, fmt, delimiter, records_path, xml_record_tag, sheet
        )
